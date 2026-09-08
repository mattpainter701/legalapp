"""Create synthetic, independently extractable live retrieval fixtures."""

from pathlib import Path
import hashlib
import json
from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from pypdf import PdfReader

OUT = Path(__file__).resolve().parents[1] / "output" / "rag-validation"
OUT.mkdir(parents=True, exist_ok=True)
fixtures = {
    "CSA_RAG_20260907_Kestrel_Intake.txt": (
        "Kestrel validation intake",
        [
            "Synthetic retrieval test record. This is fictional and contains no client information.",
            "Record family: KESTREL-7294. Revision date: September 7, 2026.",
            "The Kestrel project archive access phrase is indigo otter lantern 7294.",
            "The designated archive custodian is Mira Quill.",
            "The intake record lists the review meeting as October 14, 2026. A later signed amendment may supersede that date.",
            "The intake record does not state a bank account number or a settlement amount.",
        ],
    ),
    "CSA_RAG_20260907_Kestrel_Inventory.pdf": (
        "Kestrel validation inventory",
        [
            "Synthetic retrieval test record. This is fictional and contains no client information.",
            "Record family: KESTREL-7294. Inventory certified September 8, 2026.",
            "The sealed archive contains exactly 37 amber folders and 12 violet folders.",
            "The archive location is Cabinet Q7 on the second floor.",
            "The inventory verification code is COPPER-FINCH-5831.",
            "The custodian must reconcile the folder counts before the review meeting.",
        ],
    ),
    "CSA_RAG_20260907_Kestrel_Amendment.docx": (
        "Kestrel validation amendment",
        [
            "Synthetic retrieval test record. This is fictional and contains no client information.",
            "Record family: KESTREL-7294. Signed amendment dated September 9, 2026.",
            "This amendment supersedes only the review meeting date in the September 7 intake record. The review meeting is now October 21, 2026 at 10:30 AM Central Time.",
            "The amendment approver is Rowan Vale. The approval marker is SILVER-BADGER-9062.",
            "All archive access, custodian, and inventory details remain unchanged.",
            "The amendment contains no settlement amount and does not authorize any payment.",
        ],
    ),
}
manifest = {}
for name, (title, paragraphs) in fixtures.items():
    path = OUT / name
    if path.suffix == ".txt":
        path.write_text(
            title + "\n\n" + "\n\n".join(paragraphs) + "\n", encoding="utf-8"
        )
        extracted = path.read_text(encoding="utf-8")
    elif path.suffix == ".docx":
        doc = Document()
        doc.sections[0].top_margin = Inches(0.8)
        doc.sections[0].bottom_margin = Inches(0.8)
        doc.styles["Normal"].font.size = Pt(11)
        doc.add_heading(title, 0)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
        doc.save(path)
        extracted = "\n".join(p.text for p in Document(path).paragraphs)
    else:
        styles = getSampleStyleSheet()
        story = [Paragraph(title, styles["Title"]), Spacer(1, 16)]
        for paragraph in paragraphs:
            story.extend([Paragraph(paragraph, styles["BodyText"]), Spacer(1, 12)])
        SimpleDocTemplate(str(path), topMargin=58, bottomMargin=58).build(story)
        extracted = "\n".join(p.extract_text() for p in PdfReader(path).pages)
    assert all(paragraph in extracted for paragraph in paragraphs), name
    manifest[name] = {
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "extracted_text": extracted,
    }
(OUT / "answer-manifest.json").write_text(
    json.dumps(manifest, indent=2), encoding="utf-8"
)
print(json.dumps({"output": str(OUT), "files": list(manifest)}, indent=2))
