"""Pushing authored documents and templates in through Workspace MCP.

The two capabilities under test are the write half of the MCP product: an
outside agent that built a DOCX or authored a firm template can hand it to
LawHand.  Both must land as review work — never as approved output — so these
tests assert the draft/inactive and preview-binding invariants as hard as they
assert the happy path.
"""

from __future__ import annotations

import base64
import hashlib
import io
import uuid
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from app.schemas.workspace_mcp import (
    MAX_DOCUMENT_BYTES,
    ProposeDocumentTemplateArgs,
    ProposeMatterDocumentFileArgs,
    ProposeMatterFileArgs,
)
from app.services import document_template_push as push
from app.services.automation_capabilities import (
    ApprovalPolicy,
    CapabilityContext,
    CapabilityEffect,
    CapabilityError,
    capability_catalog,
    resolve_capability_spec,
)
from app.services.chat_tools import handlers
from app.services.document_template_versions import body_sha256
from app.services.matter_file_push import resolve_pushed_file


class _DB:
    """Session double resolving the two query shapes the service issues.

    Templates are looked up by the bound primary key, and the version counter
    lookup (only reached for a template that predates versioning) returns none.
    """

    def __init__(self, templates=()):
        self.templates = {template.id: template for template in templates}
        self.added = []
        self.flushes = 0

    async def scalar(self, statement):
        if "max(" in str(statement).lower():
            return None
        bound = set(statement.compile().params.values())
        for template_id, template in self.templates.items():
            if template_id in bound or str(template_id) in bound:
                return template
        return None

    async def flush(self):
        self.flushes += 1

    def add(self, instance):
        self.added.append(instance)
        template_id = getattr(instance, "id", None)
        if template_id is not None and hasattr(instance, "body"):
            self.templates[template_id] = instance


TENANT_ID = uuid.uuid4()
USER_ID = uuid.uuid4()


@pytest.fixture
def upload_dir(tmp_path, monkeypatch):
    """Point retained template sources at a per-test directory.

    Settings are process-wide, so this has to be undone after each test rather
    than assigned in place.
    """

    from app.services import document_template_push

    monkeypatch.setattr(
        document_template_push.settings, "UPLOAD_DIR", str(tmp_path), raising=False
    )
    return tmp_path


def _context(db, *, idempotency_key=None):
    return CapabilityContext(
        db=db,
        user=SimpleNamespace(id=USER_ID, tenant_id=TENANT_ID),
        channel="workspace_mcp",
        granted_scopes=frozenset(
            {
                "matters:read",
                "templates:read",
                "templates:propose",
                "documents:propose",
            }
        ),
        client_id="claude-desktop",
        grant_id=uuid.uuid4(),
        idempotency_key=idempotency_key,
    )


def _template_row(**overrides):
    row = SimpleNamespace(
        id=uuid.uuid4(),
        tenant_id=TENANT_ID,
        title="Engagement Letter",
        body="Dear {{client_name}}",
        category="engagement_letter",
        description=None,
        module="civil",
        stage="intake",
        jurisdiction="North Dakota",
        kind="letter",
        variable_schema=None,
        format="markdown",
        status="draft",
        is_active=False,
        visibility="tenant",
        current_version_no=1,
        source_provenance=None,
        source_storage_path=None,
        source_sha256=None,
        source_filename=None,
    )
    for key, value in overrides.items():
        setattr(row, key, value)
    return row


def _args(**overrides):
    payload = {
        "title": "Engagement Letter",
        "body": "Dear {{client_name}},\n\nWe represent you in {{matter_name}}.",
        "category": "engagement_letter",
    }
    payload.update(overrides)
    return ProposeDocumentTemplateArgs(**payload)


def _docx_bytes(paragraphs=("Settlement Agreement", "The parties agree.")):
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── Capability contract ─────────────────────────────────────────────────────


def test_push_capabilities_are_workspace_only_review_work():
    workspace = {item["name"] for item in capability_catalog(audience="workspace_mcp")}
    chat = {item["name"] for item in capability_catalog(audience="matter_chat")}
    for name in ("propose_matter_document_file", "propose_document_template"):
        assert name in workspace
        # Matter chat authors bodies in LawHand; only external clients push files.
        assert name not in chat
        spec = resolve_capability_spec(name)
        assert spec.effect == CapabilityEffect.PROPOSE
        assert spec.approval_policy == ApprovalPolicy.LAWHAND_REVIEW
        assert spec.mcp_annotations()["destructiveHint"] is False
        assert spec.mcp_annotations()["readOnlyHint"] is False

    assert resolve_capability_spec("propose_document_template").required_scopes == (
        "templates:read",
        "templates:propose",
    )
    assert resolve_capability_spec("propose_matter_document_file").required_scopes == (
        "matters:read",
        "documents:propose",
    )


def test_template_push_requires_its_own_consent_scope():
    spec = resolve_capability_spec("propose_document_template")
    context = _context(_DB())
    context.granted_scopes = frozenset({"templates:read"})
    with pytest.raises(CapabilityError) as excinfo:
        spec.authorize(context)
    assert excinfo.value.code == "capability_scope_denied"


# ── Template push ───────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_pushed_template_is_saved_inactive_and_versioned():
    db = _DB()
    result = await push.push_workspace_template(_context(db), _args())

    assert result["created"] is True
    assert result["is_active"] is False
    assert result["status"] == "draft"
    assert result["format"] == "markdown"
    assert result["version_no"] == 1
    assert result["variables"] == ["client_name", "matter_name"]

    template = db.added[0]
    assert template.is_active is False
    assert template.status == "draft"
    assert template.source_provenance["origin"] == "workspace_mcp"
    assert template.source_provenance["client_id"] == "claude-desktop"
    assert template.source_provenance["pushed_by_user_id"] == str(USER_ID)
    # The immutable version row is what a reviewer compares against.
    version = db.added[1]
    assert version.template_id == template.id
    assert version.is_active is False
    assert version.body_sha256 == body_sha256(template.body)


@pytest.mark.asyncio
async def test_same_client_request_id_replays_instead_of_duplicating():
    request_id = uuid.uuid4()
    db = _DB()
    first = await push.push_workspace_template(
        _context(db), _args(client_request_id=request_id)
    )
    second = await push.push_workspace_template(
        _context(db), _args(client_request_id=request_id)
    )

    assert first["template_id"] == second["template_id"]
    assert second["created"] is False
    assert second["idempotent_replay"] is True
    # One template plus one version row: the retry wrote nothing.
    assert len(db.added) == 2


@pytest.mark.asyncio
async def test_reusing_a_request_id_for_different_content_conflicts():
    request_id = uuid.uuid4()
    db = _DB()
    await push.push_workspace_template(
        _context(db), _args(client_request_id=request_id)
    )
    with pytest.raises(CapabilityError) as excinfo:
        await push.push_workspace_template(
            _context(db),
            _args(client_request_id=request_id, body="Something else entirely"),
        )
    assert excinfo.value.code == "idempotency_conflict"


@pytest.mark.asyncio
async def test_an_idempotency_header_alone_gives_a_stable_template_identity():
    db = _DB()
    first = await push.push_workspace_template(
        _context(db, idempotency_key="push-42"), _args()
    )
    second = await push.push_workspace_template(
        _context(db, idempotency_key="push-42"), _args()
    )
    assert first["template_id"] == second["template_id"]
    assert second["idempotent_replay"] is True


@pytest.mark.asyncio
async def test_a_draft_template_can_be_revised_in_place():
    existing = _template_row()
    db = _DB([existing])
    result = await push.push_workspace_template(
        _context(db),
        _args(
            template_id=existing.id,
            body="Dear {{client_name}}, revised.",
            change_summary="Tighten the fee paragraph",
        ),
    )
    assert result["created"] is False
    assert result["idempotent_replay"] is False
    assert result["version_no"] == 2
    assert existing.body == "Dear {{client_name}}, revised."
    assert existing.is_active is False
    assert db.added[0].change_summary == "Tighten the fee paragraph"


@pytest.mark.asyncio
async def test_a_live_template_is_never_edited_in_place():
    live = _template_row(is_active=True, status="approved")
    db = _DB([live])
    with pytest.raises(CapabilityError) as excinfo:
        await push.push_workspace_template(
            _context(db), _args(template_id=live.id, body="Rewritten {{client_name}}")
        )
    assert excinfo.value.code == "template_not_draft"
    assert live.body == "Dear {{client_name}}"
    assert live.is_active is True


@pytest.mark.asyncio
async def test_superseding_a_live_template_creates_a_separate_draft():
    live = _template_row(is_active=True, status="approved")
    db = _DB([live])
    result = await push.push_workspace_template(
        _context(db),
        _args(
            supersedes_template_id=live.id,
            category="engagement_letter",
            module=None,
            jurisdiction=None,
        ),
    )
    assert result["template_id"] != str(live.id)
    assert result["is_active"] is False
    assert result["supersedes_template_id"] == str(live.id)
    # The live template keeps serving the firm untouched.
    assert live.is_active is True
    assert live.body == "Dear {{client_name}}"
    # Unstated placement is inherited so the draft stays comparable.
    draft = db.added[0]
    assert draft.jurisdiction == "North Dakota"
    assert draft.module == "civil"


@pytest.mark.asyncio
async def test_a_templates_format_cannot_be_changed_in_place():
    word_template = _template_row(format="docx", source_sha256="a" * 64)
    db = _DB([word_template])
    with pytest.raises(CapabilityError) as excinfo:
        await push.push_workspace_template(
            _context(db), _args(template_id=word_template.id)
        )
    assert excinfo.value.code == "template_format_immutable"


@pytest.mark.asyncio
async def test_unknown_category_is_refused():
    with pytest.raises(CapabilityError) as excinfo:
        await push.push_workspace_template(
            _DB(), _args(category="not_a_category")
        )
    assert excinfo.value.code == "invalid_template_category"


@pytest.mark.asyncio
async def test_a_field_map_may_not_name_variables_the_body_lacks():
    with pytest.raises(CapabilityError) as excinfo:
        await push.push_workspace_template(
            _DB(),
            _args(variable_schema={"fields": [{"name": "opposing_counsel"}]}),
        )
    assert excinfo.value.code == "invalid_variable_schema"
    assert "opposing_counsel" in excinfo.value.message


def test_logic_markers_are_not_reported_as_variables():
    variables = push.template_variable_names(
        "{{#if fee_shifting}}{{fee_clause}}{{/if}} {{client_name}}"
    )
    assert variables == ["fee_clause", "client_name"]


def test_template_id_and_supersedes_are_mutually_exclusive():
    with pytest.raises(ValueError):
        ProposeDocumentTemplateArgs(
            title="x",
            body="y",
            template_id=uuid.uuid4(),
            supersedes_template_id=uuid.uuid4(),
        )


# ── Document push ───────────────────────────────────────────────────────────


def _file_args(**overrides):
    payload = {
        "matter_id": uuid.uuid4(),
        "title": "Settlement Agreement",
        "content_base64": base64.b64encode(_docx_bytes()).decode(),
    }
    payload.update(overrides)
    return ProposeMatterDocumentFileArgs(**payload)


@pytest.mark.asyncio
async def test_pushed_docx_binds_review_text_to_the_uploaded_bytes(monkeypatch):
    source = _docx_bytes()
    captured = {}

    async def fake_propose(_context, proposal, **kwargs):
        captured["proposal"] = proposal
        captured["kwargs"] = kwargs
        return {"task_id": str(uuid.uuid4()), "status": "review"}

    monkeypatch.setattr(handlers, "_propose_matter_document", fake_propose)
    result = await handlers.propose_matter_document_file(
        _context(_DB()),
        _file_args(content_base64=base64.b64encode(source).decode()),
    )

    # The caller never supplies review text; it is extracted from the file.
    assert captured["proposal"].body == "Settlement Agreement\nThe parties agree."
    assert captured["kwargs"]["source_docx_bytes"] == source
    assert result["uploaded_size"] == len(source)
    assert len(result["uploaded_sha256"]) == 64


@pytest.mark.asyncio
async def test_a_mismatched_content_digest_is_refused(monkeypatch):
    async def fail(*_args, **_kwargs):  # pragma: no cover - must not run
        raise AssertionError("A document that failed its digest must not be proposed")

    monkeypatch.setattr(handlers, "_propose_matter_document", fail)
    with pytest.raises(CapabilityError) as excinfo:
        await handlers.propose_matter_document_file(
            _context(_DB()), _file_args(content_sha256="b" * 64)
        )
    assert excinfo.value.code == "document_integrity_failed"


@pytest.mark.asyncio
async def test_a_file_that_is_not_a_docx_is_refused():
    with pytest.raises(CapabilityError) as excinfo:
        await handlers.propose_matter_document_file(
            _context(_DB()),
            _file_args(content_base64=base64.b64encode(b"%PDF-1.7 not a docx").decode()),
        )
    assert excinfo.value.code == "invalid_docx_package"


@pytest.mark.asyncio
async def test_macro_enabled_content_is_refused():
    import zipfile

    buffer = io.BytesIO(_docx_bytes())
    with zipfile.ZipFile(buffer, "a") as archive:
        archive.writestr("word/vbaProject.bin", b"macro payload")
    with pytest.raises(CapabilityError) as excinfo:
        await handlers.propose_matter_document_file(
            _context(_DB()),
            _file_args(content_base64=base64.b64encode(buffer.getvalue()).decode()),
        )
    assert excinfo.value.code == "active_or_embedded_content"


@pytest.mark.asyncio
async def test_malformed_base64_is_a_named_argument_error():
    args = _file_args().model_copy(update={"content_base64": "not base64 !!!"})
    with pytest.raises(CapabilityError) as excinfo:
        await handlers.propose_matter_document_file(_context(_DB()), args)
    assert excinfo.value.code == "invalid_document_encoding"


def test_the_pushed_file_size_cap_is_enforced_before_decoding():
    from pydantic import ValidationError

    oversized = "A" * (((MAX_DOCUMENT_BYTES + 2) // 3) * 4 + 4)
    with pytest.raises(ValidationError):
        _file_args(content_base64=oversized)


def test_a_pushed_filename_may_not_carry_a_path():
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        _file_args(filename="../../etc/passwd.docx")
    with pytest.raises(ValidationError):
        _file_args(filename="agreement.pdf")


# ── Source-backed templates: DOCX and fillable PDF ──────────────────────────


def _fillable_pdf(fields=("client_name", "matter_number")):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    page = canvas.Canvas(buffer, pagesize=letter)
    page.drawString(72, 740, "Client Intake Form")
    form = page.acroForm
    for index, name in enumerate(fields):
        form.textfield(
            name=name, tooltip=name, x=72, y=700 - index * 30, width=200, height=20
        )
    page.save()
    return buffer.getvalue()


def _flat_pdf():
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    page = canvas.Canvas(buffer, pagesize=letter)
    for index, line in enumerate(
        [
            "RETAINER AGREEMENT",
            "This agreement is made between the firm and the client.",
            "The client agrees to the attached fee schedule.",
        ]
    ):
        page.drawString(72, 720 - index * 24, line)
    page.save()
    return buffer.getvalue()


def _word_template_bytes():
    document = Document()
    document.add_heading("Engagement Letter", 0)
    document.add_paragraph("Dear Jane Smith,")
    document.add_paragraph("We are pleased to represent you in this matter.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _source_args(raw, *, fmt, filename, **overrides):
    payload = {
        "title": "Client Intake Form",
        "format": fmt,
        "category": "other",
        "content_base64": base64.b64encode(raw).decode(),
        "filename": filename,
    }
    payload.update(overrides)
    return ProposeDocumentTemplateArgs(**payload)


@pytest.mark.asyncio
async def test_a_fillable_pdf_template_arrives_complete_and_inactive(
    upload_dir,
):
    db = _DB()
    result = await push.push_workspace_template(
        _context(db), _source_args(_fillable_pdf(), fmt="pdf", filename="intake.pdf")
    )

    assert result["format"] == "pdf"
    assert result["is_active"] is False
    assert result["status"] == "draft"
    # The field map came from the form itself, not from the caller.
    assert result["fillable_field_count"] == 2
    assert result["variables"] == ["client_name", "matter_number"]

    template = db.added[0]
    assert template.source_sha256 == result["source_sha256"]
    # The retained source is what later renders, so it must be on disk and match.
    stored = Path(template.source_storage_path)
    assert stored.is_file()
    assert hashlib.sha256(stored.read_bytes()).hexdigest() == template.source_sha256
    fields = template.variable_schema["fields"]
    assert {field["pdf_field_name"] for field in fields} == {
        "client_name",
        "matter_number",
    }


@pytest.mark.asyncio
async def test_a_flat_pdf_is_sent_to_the_review_canvas_instead(upload_dir):
    with pytest.raises(CapabilityError) as excinfo:
        await push.push_workspace_template(
            _context(_DB()),
            _source_args(_flat_pdf(), fmt="pdf", filename="retainer.pdf"),
        )
    assert excinfo.value.code == "pdf_not_fillable"
    # A refusal has to say where the work can actually be done.
    assert "canvas" in excinfo.value.message


@pytest.mark.asyncio
async def test_a_word_template_keeps_its_source_and_discovered_anchors(
    upload_dir,
):
    db = _DB()
    result = await push.push_workspace_template(
        _context(db),
        _source_args(
            _word_template_bytes(),
            fmt="docx",
            filename="engagement.docx",
            title="Engagement Letter",
            category="engagement_letter",
        ),
    )
    assert result["format"] == "docx"
    assert result["is_active"] is False
    template = db.added[0]
    assert Path(template.source_storage_path).is_file()
    # Every discovered Word field names the exact source text it replaces.
    for field in template.variable_schema["fields"]:
        assert str(field.get("source_text") or "").strip()


@pytest.mark.asyncio
async def test_declared_format_cannot_disagree_with_the_uploaded_bytes(upload_dir):
    with pytest.raises(CapabilityError) as excinfo:
        await push.push_workspace_template(
            _context(_DB()),
            _source_args(_fillable_pdf(), fmt="docx", filename="intake.docx"),
        )
    assert excinfo.value.code in {
        "template_format_mismatch",
        "template_analysis_failed",
    }


@pytest.mark.asyncio
async def test_a_mismatched_template_digest_is_refused():
    with pytest.raises(CapabilityError) as excinfo:
        await push.push_workspace_template(
            _context(_DB()),
            _source_args(
                _fillable_pdf(),
                fmt="pdf",
                filename="intake.pdf",
                content_sha256="c" * 64,
            ),
        )
    assert excinfo.value.code == "template_integrity_failed"


def test_a_markdown_template_may_not_carry_a_file():
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        ProposeDocumentTemplateArgs(
            title="t", body="b", content_base64=base64.b64encode(b"x").decode()
        )
    with pytest.raises(ValidationError):
        ProposeDocumentTemplateArgs(title="t", format="pdf")


# ── Matter file artifacts ───────────────────────────────────────────────────


def _png_bytes():
    import struct
    import zlib

    def chunk(tag, data):
        payload = tag + data
        return (
            struct.pack(">I", len(data))
            + payload
            + struct.pack(">I", zlib.crc32(payload))
        )

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(b"\x00\x00\x00\x00"))
        + chunk(b"IEND", b"")
    )


_EML_BYTES = (
    b"From: clerk@court.example\r\nTo: firm@example.com\r\n"
    b"Subject: Filing confirmation\r\n\r\nYour filing was accepted.\r\n"
)


def _file_push_args(filename, raw, **overrides):
    payload = {
        "matter_id": uuid.uuid4(),
        "filename": filename,
        "content_base64": base64.b64encode(raw).decode(),
    }
    payload.update(overrides)
    return ProposeMatterFileArgs(**payload)


@pytest.mark.parametrize(
    ("filename", "raw", "content_type"),
    [
        ("exhibit-a.png", _png_bytes(), "image/png"),
        ("confirmation.eml", _EML_BYTES, "message/rfc822"),
        ("timeline.csv", b"date,event\n2026-01-01,filed\n", "text/csv"),
        ("notes.txt", b"call summary", "text/plain"),
    ],
)
def test_supported_matter_artifacts_resolve_to_their_real_type(
    filename, raw, content_type
):
    resolved = resolve_pushed_file(_file_push_args(filename, raw))
    assert resolved.content_type == content_type
    assert resolved.sha256 == hashlib.sha256(raw).hexdigest()
    assert resolved.filename == filename


@pytest.mark.parametrize(
    ("label", "filename", "raw", "code"),
    [
        ("executable renamed", "photo.png", b"MZ\x90\x00 payload", "file_content_mismatch"),
        ("archive", "bundle.zip", b"PK\x03\x04", "unsupported_file_type"),
        (
            "legacy macro container",
            "report.docx",
            bytes.fromhex("D0CF11E0A1B11AE1") + b"payload",
            "file_content_mismatch",
        ),
    ],
)
def test_matter_artifacts_that_lie_about_their_bytes_are_refused(
    label, filename, raw, code
):
    with pytest.raises(CapabilityError) as excinfo:
        resolve_pushed_file(_file_push_args(filename, raw))
    assert excinfo.value.code == code, label


def test_a_matter_artifact_digest_is_verified():
    with pytest.raises(CapabilityError) as excinfo:
        resolve_pushed_file(
            _file_push_args("exhibit-a.png", _png_bytes(), content_sha256="d" * 64)
        )
    assert excinfo.value.code == "file_integrity_failed"


def test_a_matter_artifact_needs_a_real_extension():
    from pydantic import ValidationError

    for bad in ("noextension", ".hidden", "../escape.png"):
        with pytest.raises(ValidationError):
            _file_push_args(bad, _png_bytes())


def test_matter_file_capability_is_workspace_only_review_work():
    spec = resolve_capability_spec("propose_matter_file")
    assert spec.audiences == ("workspace_mcp",)
    assert spec.effect == CapabilityEffect.PROPOSE
    assert spec.approval_policy == ApprovalPolicy.LAWHAND_REVIEW
    assert spec.required_scopes == ("matters:read", "documents:propose")
