import copy
import io
from types import SimpleNamespace
from unittest.mock import AsyncMock
import uuid

import pytest
from docx import Document
from fastapi import HTTPException

from app.routers import document_templates as router
from app.schemas.document_template import DocumentTemplateWordCleanupRequest
from app.services.docx_placeholder_authoring import cleanup_docx_source, schema_after_word_edit
from app.services.docx_templates import TemplateDocxError, docx_source_key, fill_docx_template


def field(name, start, end, ordinal=0, **extra):
    anchor = {"paragraph_ordinal": ordinal, "start": start, "end": end}
    return {"name": name, "source_text": "Alex", "docx_anchor": anchor,
            "docx_source_key": docx_source_key("Alex", anchor), **extra}


def test_wording_edit_preserves_format_and_rebases_a_later_field_for_generation():
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Dear ").italic = True
    paragraph.add_run("Alex").bold = True
    paragraph.add_run(", welcome.")
    stream = io.BytesIO()
    doc.save(stream)
    schema = {"version": 2, "fields": [field("client", 5, 9, binding="client.name")], "source_review": {"old": "fixed"}}
    original_schema = copy.deepcopy(schema)
    edits = dict(paragraph_ordinal=0, start=0, end=4, replacement_text="Hello there")
    updated = cleanup_docx_source(stream.getvalue(), original_text="Dear", **edits)
    mapped = schema_after_word_edit(schema, **edits)
    assert mapped["fields"][0]["docx_anchor"] == {"paragraph_ordinal": 0, "start": 12, "end": 16}
    assert mapped["fields"][0]["docx_source_key"] != schema["fields"][0]["docx_source_key"]
    assert mapped["fields"][0]["binding"] == "client.name"
    assert mapped["source_review"] == {} and schema == original_schema
    output = fill_docx_template(updated, variable_schema=mapped, variables={"client": "Taylor"})
    result = Document(io.BytesIO(output))
    assert result.paragraphs[0].text == "Hello there Taylor, welcome."
    assert result.paragraphs[0].runs[0].italic
    assert any(run.text == "Taylor" and run.bold for run in result.paragraphs[0].runs)


def test_edit_preserves_other_paragraphs_earlier_anchors_and_unanchored_fields():
    fields = [field("earlier", 0, 4), field("other", 5, 9, ordinal=1), {"name": "token"}, field("no_key", 20, 24)]
    fields[-1].pop("docx_source_key")
    schema = schema_after_word_edit({"fields": fields}, paragraph_ordinal=0, start=10, end=12, replacement_text="X")
    assert schema["fields"][:3] == fields[:3]
    assert schema["fields"][-1]["docx_anchor"]["start"] == 19


def test_overlapping_included_field_is_rejected_and_excluded_mapping_is_retired():
    edit = dict(paragraph_ordinal=0, start=5, end=9, replacement_text="Taylor")
    with pytest.raises(TemplateDocxError, match="mapped field"):
        schema_after_word_edit({"fields": [field("client", 5, 9)]}, **edit)
    assert schema_after_word_edit({"fields": [field("client", 5, 9, included=False)]}, **edit)["fields"] == []
    assert schema_after_word_edit(None, **edit) == {"fields": [], "source_review": {}}


@pytest.mark.asyncio
@pytest.mark.parametrize("expected", [{"expected_source_sha256": "b" * 64}, {"expected_version_no": 3}])
async def test_stale_wording_request_does_not_read_or_create_source(monkeypatch, expected):
    template = SimpleNamespace(format="docx", source_sha256="a" * 64, current_version_no=4)
    db = SimpleNamespace(scalar=AsyncMock(return_value=template))
    monkeypatch.setattr(router, "set_tenant_context", AsyncMock())
    source = AsyncMock()
    monkeypatch.setattr(router, "_verified_template_source", source)
    payload = DocumentTemplateWordCleanupRequest(paragraph_ordinal=0, start=0, end=1, original_text="A", replacement_text="B", **expected)
    with pytest.raises(HTTPException) as failure:
        await router.cleanup_word_draft(uuid.uuid4(), payload, SimpleNamespace(tenant_id=uuid.uuid4()), db)
    assert failure.value.status_code == 409
    source.assert_not_awaited()
