from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from app.services.docx_templates import TemplateDocxError, fill_docx_template
from app.services.esign.placement import (
    generated_signing_metadata,
    is_signing_template_field,
)
from app.services.pdf_templates import (
    TemplatePdfError,
    discover_pdf_fields,
    fill_pdf_template,
    pdf_review_evidence,
    validate_representative_pdf_variables,
)


def pdf_fixture(form, mixed=False):
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=(612, 792))
    pdf.drawString(72, 740, "Signing date test")
    if form:
        for index, name in enumerate(["signed_on", "event_date"]):
            pdf.acroForm.textfield(
                name=name,
                value="OLD SAMPLE",
                x=72,
                y=600 - index * 50,
                width=200,
                height=24,
            )
    pdf.showPage()
    pdf.save()
    content = output.getvalue()
    fields = (
        discover_pdf_fields(content)
        if form
        else [
            {
                "name": name,
                "pdf_source_key": name,
                "pdf_overlay": {
                    "page": 1,
                    "rect": [72, 600 - index * 50, 272, 624 - index * 50],
                    "source_kind": "manual",
                },
            }
            for index, name in enumerate(["signed_on", "event_date"])
        ]
    )
    for field in fields:
        field.update(field_type="date", required=True, included=True)
        if field["name"] == "signed_on":
            field["signer_role"] = "client"
    if mixed:
        fields.append(
            {
                "name": "second_signing_date",
                "field_type": "date",
                "signer_role": "attorney",
                "required": True,
                "pdf_source_key": "second-date",
                "pdf_overlay": {
                    "page": 1,
                    "rect": [72, 400, 272, 424],
                    "source_kind": "manual",
                },
            }
        )
    return content, {"fields": fields}


@pytest.mark.parametrize(
    "form,flatten,mixed",
    [
        (False, True, False),
        (True, True, False),
        (True, False, False),
        (True, True, True),
    ],
)
def test_signing_date_does_not_need_a_fill_value_or_retain_sample_form_data(
    form, flatten, mixed
):
    content, schema = pdf_fixture(form, mixed)
    values = {"event_date": "2026-10-12"}
    validate_representative_pdf_variables(schema, values)
    assert pdf_review_evidence(schema, values) == (["event_date"], 1)
    output = fill_pdf_template(
        content,
        variable_schema=schema,
        variables=values,
        enforce_required=True,
        flatten=flatten,
    )
    document = PdfReader(BytesIO(output))
    if flatten:
        assert "2026-10-12" in document.pages[0].extract_text()
        assert "OLD SAMPLE" not in document.pages[0].extract_text()
    else:
        assert document.get_fields()["signed_on"]["/V"] == ""
        assert document.get_fields()["event_date"]["/V"] == "2026-10-12"
    with pytest.raises(TemplatePdfError, match="Unknown PDF template variable"):
        fill_pdf_template(
            content,
            variable_schema=schema,
            variables={**values, "signed_on": "2026-01-01"},
            flatten=flatten,
        )
    with pytest.raises(TemplatePdfError, match="representative"):
        validate_representative_pdf_variables(schema, {})


def test_required_word_signatures_and_signing_dates_stay_blank_and_keep_signer_roles():
    doc = Document()
    doc.add_paragraph(
        "Signature: {{signature}} Signed: {{signed_on}} Event: {{event_date}}"
    )
    stream = BytesIO()
    doc.save(stream)
    fields = [
        {
            "name": "signature",
            "field_type": "signature",
            "signer_role": "client",
            "required": True,
        },
        {
            "name": "signed_on",
            "field_type": "date",
            "signer_role": "client",
            "required": True,
        },
        {"name": "event_date", "field_type": "date", "required": True},
    ]
    schema = {"fields": fields}
    values = {"event_date": "2026-10-12"}
    output = fill_docx_template(
        stream.getvalue(),
        variable_schema=schema,
        variables=values,
        enforce_required=True,
    )
    assert (
        Document(BytesIO(output)).paragraphs[0].text
        == "Signature:  Signed:  Event: 2026-10-12"
    )
    assert values == {"event_date": "2026-10-12"}
    assert generated_signing_metadata(
        schema, source=output, template_format="docx"
    ) == ([], ["client"], True)
    for name in ["signature", "signed_on"]:
        with pytest.raises(TemplateDocxError, match="must remain blank"):
            fill_docx_template(
                stream.getvalue(),
                variable_schema=schema,
                variables={**values, name: "prefilled"},
            )
    with pytest.raises(TemplateDocxError, match="event_date"):
        fill_docx_template(
            stream.getvalue(),
            variable_schema=schema,
            variables={},
            enforce_required=True,
        )
    assert not is_signing_template_field({"field_type": "date", "signer_role": "  "})
    assert not is_signing_template_field(None)
