"""Synthetic regressions for customer Word master preparation; no customer data."""

import io
import json
import zipfile
import uuid
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi import HTTPException

from app.services.docx_outline import docx_outline, validate_visual_field_map
from app.services.docx_templates import (
    TemplateDocxError,
    fill_docx_template,
    iter_docx_paragraphs,
)
from app.services.docx_source_review import (
    candidate_is_mapped,
    require_source_review,
    source_review_candidates,
    validate_review_decisions,
    validate_value_links,
    word_values,
)
from app.services.template_intake import _suggest_docx_template, analyze_template_upload


def source(*texts, build=None):
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    if build:
        build(doc)
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def discover(content):
    text = "\n".join(
        p.text for p in iter_docx_paragraphs(Document(io.BytesIO(content)))
    )
    _, fields, warnings = _suggest_docx_template(content, text)
    return [field.as_dict() for field in fields], warnings


def filled_text(content, fields, values):
    output = fill_docx_template(
        content, variable_schema={"fields": fields}, variables=values
    )
    return "\n".join(p.text for p in iter_docx_paragraphs(Document(io.BytesIO(output))))


def test_generic_repeated_tokens_are_independent_and_explicit_links_share_values():
    content = source(
        "Car value: [AMOUNT]; loan: [AMOUNT]", "[HUSBAND NAME] / [HUSBAND NAME]"
    )
    fields, warnings = discover(content)
    amounts = [field for field in fields if field["source_text"] == "[AMOUNT]"]
    assert len(amounts) == 2
    assert all(field["docx_anchor"] and field["context"] for field in amounts)
    assert (
        len([field for field in fields if field["source_text"] == "[HUSBAND NAME]"])
        == 1
    )
    values = {amounts[0]["name"]: "$900", amounts[1]["name"]: "$200"}
    assert "Car value: $900; loan: $200" in filled_text(content, fields, values)
    amounts[1]["value_from"] = amounts[0]["name"]
    assert "Car value: $900; loan: $900" in filled_text(content, fields, values)
    assert any("separated by location" in warning for warning in warnings)


def test_labelled_identity_keeps_definition_and_updates_signature():
    content = source(
        "Client: Taylor Example (hereinafter “Client”)", "CLIENT: Taylor Example"
    )
    fields, _ = discover(content)
    client = next(field for field in fields if field["name"] == "client_name")
    assert client["source_text"] == "Taylor Example"
    result = filled_text(content, fields, {"client_name": "Jordan Sample"})
    assert "Client: Jordan Sample (hereinafter “Client”)" in result
    assert "CLIENT: Jordan Sample" in result
    assert "Taylor Example" not in result


def test_leading_choice_marks_belong_to_following_option_and_reject_double_answer():
    content = source(
        "Who keeps the car? ___ Husband ___ Wife",
        "Is insurance current?",
        "___ yes",
        "___ no",
    )
    fields, _ = discover(content)
    assert [field["docx_choice"]["option"] for field in fields] == [
        "Husband",
        "Wife",
        "Yes",
        "No",
    ]
    assert all(field["field_type"] == "checkbox" for field in fields)
    values = {fields[1]["name"]: "true", fields[2]["name"]: "true"}
    assert "Who keeps the car?  Husband X Wife" in filled_text(content, fields, values)
    values[fields[0]["name"]] = "true"
    with pytest.raises(TemplateDocxError, match="only one"):
        filled_text(content, fields, values)
    with pytest.raises(TemplateDocxError, match="checked or unchecked"):
        word_values(fields, {fields[0]["name"]: "maybe"})


def test_explicit_check_all_allows_multiple_answers():
    fields, _ = discover(source("Check all owners: ___ Husband ___ Wife"))
    assert not fields[0]["docx_choice"]["exclusive"]
    assert list(
        word_values(fields, {field["name"]: "true" for field in fields}).values()
    ) == ["X", "X"]


def test_numbered_questions_produce_valid_field_names():
    content = source("1. Which party pays? ___ Husband ___ Wife")
    fields, _ = discover(content)
    validate_visual_field_map(content, {"fields": fields}, {"fields": fields})
    assert all(field["name"].startswith("answer_1_") for field in fields)


@pytest.mark.parametrize(
    "fmt, expected",
    [
        ("lowerLetter", "b."),
        ("upperLetter", "B."),
        ("lowerRoman", "ii."),
        ("upperRoman", "II."),
        ("decimal", "2."),
    ],
)
def test_numbering_formats_and_start_override(fmt, expected):
    def build(doc):
        paragraph = doc.add_paragraph("Numbered", style="List Number")
        numid = paragraph.style.element.pPr.numPr.numId.val
        numbering = doc.part.numbering_part.element
        num = next(
            node
            for node in numbering.findall(qn("w:num"))
            if int(node.get(qn("w:numId"))) == numid
        )
        abstractid = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        abstract = next(
            node
            for node in numbering.findall(qn("w:abstractNum"))
            if node.get(qn("w:abstractNumId")) == abstractid
        )
        abstract.find(qn("w:lvl")).find(qn("w:numFmt")).set(qn("w:val"), fmt)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), "2")
        override.append(start)
        num.append(override)

    assert docx_outline(source(build=build))["paragraphs"][0]["numbering"] == expected


def test_review_can_exclude_an_existing_global_field_without_reanchoring():
    content = source("Client: [NAME]")
    fields, _ = discover(content)
    validate_visual_field_map(
        content, {"fields": fields}, {"fields": [{**fields[0], "included": False}]}
    )
    assert "[NAME]" in filled_text(content, [{**fields[0], "included": False}], {})


@pytest.mark.parametrize(
    "patch",
    [
        {"value_from": "missing"},
        {"value_from": "a"},
        {"value_from": []},
        {"docx_choice": {}},
        {"docx_choice": "yes"},
    ],
)
def test_invalid_link_or_choice_is_rejected(patch):
    with pytest.raises(TemplateDocxError):
        validate_value_links([{"name": "a", **patch}])


@pytest.mark.parametrize(
    "target",
    [
        {"name": "b", "included": False},
        {"name": "b", "value_from": "c"},
        {"name": "b", "field_type": "checkbox"},
    ],
)
def test_links_require_an_independent_included_compatible_target(target):
    with pytest.raises(TemplateDocxError):
        validate_value_links([{"name": "a", "value_from": "b"}, target, {"name": "c"}])


def test_malformed_fields_raise_a_customer_error():
    with pytest.raises(TemplateDocxError):
        validate_value_links([None])


def test_semantic_validation_reports_invalid_links():
    from app.services.template_semantics import (
        validate_semantic_metadata,
        TemplateSemanticsError,
    )

    with pytest.raises(TemplateSemanticsError, match="independent field"):
        validate_semantic_metadata({"fields": [{"name": "a", "value_from": "missing"}]})


def test_noop_fill_preserves_every_xml_part_and_missing_header():
    def build(doc):
        paragraph = doc.add_paragraph("Client: ")
        paragraph.add_run("Taylor ").bold = True
        paragraph.add_run("Example").italic = True
        doc.sections[0].footer.paragraphs[0].text = "Footer"

    content = source(build=build)
    fields, _ = discover(content)
    output = fill_docx_template(
        content,
        variable_schema={"fields": fields},
        variables={field["name"]: field["source_text"] for field in fields},
    )
    with (
        zipfile.ZipFile(io.BytesIO(content)) as before,
        zipfile.ZipFile(io.BytesIO(output)) as after,
    ):
        assert set(before.namelist()) == set(after.namelist())
        for name in before.namelist():
            assert before.read(name) == after.read(name), name
    assert "Header" not in [p["text"] for p in docx_outline(content)["paragraphs"]]


def test_outline_visiting_missing_stories_does_not_create_parts_and_keeps_first_even_stories():
    def build(doc):
        doc.sections[0].first_page_header.paragraphs[0].text = "First [NAME]"
        doc.sections[0].even_page_footer.paragraphs[0].text = "Even [NAME]"
        doc.add_section()

    content = source("Body", build=build)
    doc = Document(io.BytesIO(content))
    before = doc.element.xml
    paragraphs = list(iter_docx_paragraphs(doc))
    assert doc.element.xml == before
    assert [p.text for p in paragraphs].count("First [NAME]") == 1
    assert [p.text for p in paragraphs].count("Even [NAME]") == 1
    assert "First Sam" in filled_text(
        content, [{"name": "name", "source_text": "[NAME]"}], {"name": "Sam"}
    )


def test_source_order_tables_merges_and_numbering_are_displayed_without_changing_ordinals():
    def build(doc):
        doc.add_paragraph("First", style="List Number")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(1, 0)).text = "Vertical"
        table.cell(0, 1).text = "Cell"
        table.cell(1, 1).add_table(rows=1, cols=1).cell(0, 0).text = "Nested"
        doc.add_paragraph("Second", style="List Number")
        wide = doc.add_table(rows=1, cols=2)
        wide.cell(0, 0).merge(wide.cell(0, 1)).text = "Wide"
        control = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        paragraph = doc.add_paragraph("Control")
        content.append(paragraph._p)
        control.append(content)
        doc.element.body.append(control)

    outline = docx_outline(source(build=build))
    assert [block["kind"] for block in outline["blocks"]] == [
        "paragraph",
        "table",
        "paragraph",
        "table",
        "paragraph",
    ]
    assert outline["blocks"][1]["rows"][0]["cells"][0]["rowspan"] == 2
    assert outline["blocks"][3]["rows"][0]["cells"][0]["colspan"] == 2
    assert [p["numbering"] for p in outline["paragraphs"] if "numbering" in p] == [
        "1.",
        "2.",
    ]
    assert outline["blocks"][2]["ordinal"] == 1  # legacy body-first anchors


def test_candidates_dispositions_and_new_master_publication_gate():
    content = source("Client: Taylor Example", "Fee: $400", "Signed: ______")
    outline = docx_outline(content)
    assert {item["kind"] for item in outline["review_candidates"]} == {
        "identity",
        "value",
        "blank",
    }
    schema = {"fields": [], "source_review_version": 1}
    with pytest.raises(TemplateDocxError, match="before publishing"):
        require_source_review(content, schema)
    decisions = {
        item["id"]: "signature" if item["kind"] == "blank" else "fixed"
        for item in outline["review_candidates"]
    }
    schema["source_review"] = decisions
    require_source_review(content, schema)
    validate_visual_field_map(
        content, {"fields": [], "source_review_version": 1}, schema
    )
    require_source_review(content, {"fields": []})  # legacy master
    with pytest.raises(TemplateDocxError, match="metadata"):
        validate_visual_field_map(
            content, schema, {**schema, "source_review_version": 0}
        )


@pytest.mark.parametrize(
    "decisions", [[], {"stale": "fixed"}, {"stale": []}, {"stale": "unknown"}]
)
def test_review_decisions_require_current_source_ids(decisions):
    with pytest.raises(TemplateDocxError):
        validate_review_decisions({"review_candidates": []}, decisions)


def test_mapped_candidate_requires_matching_placement_or_exact_global_source():
    candidate = docx_outline(source("Client: Taylor Example"))["review_candidates"][0]
    assert candidate_is_mapped(candidate, [{"source_text": "Taylor Example"}])
    assert not candidate_is_mapped(candidate, [{"source_text": "Taylor Example extra"}])
    assert not candidate_is_mapped(
        candidate, [{"source_text": "Taylor Example", "included": False}]
    )
    assert candidate_is_mapped(candidate, [{"docx_anchor": candidate["docx_anchor"]}])
    assert not candidate_is_mapped(
        candidate, [{"docx_anchor": {"paragraph_ordinal": 0, "start": 0, "end": 2}}]
    )


def test_candidate_limit_and_emphasis_do_not_silently_certify_source(monkeypatch):
    import app.services.docx_source_review as review

    monkeypatch.setattr(review, "MAX_REVIEW_CANDIDATES", 2)
    content = source("$200 $300 $400")
    outline = docx_outline(content)
    assert outline["review_truncated"]
    with pytest.raises(TemplateDocxError):
        require_source_review(
            content,
            {
                "source_review_version": 1,
                "source_review": {
                    item["id"]: "fixed" for item in outline["review_candidates"]
                },
            },
        )
    paragraphs = [
        {
            "ordinal": 0,
            "text": "Term: civil action",
            "runs": [{"start": 6, "text": "civil action", "bold": True}],
        }
    ]
    assert source_review_candidates(paragraphs)[0][0]["kind"] == "emphasized wording"
    paragraphs[0]["dynamic_field"] = True
    assert source_review_candidates(paragraphs) == ([], False)


@pytest.mark.asyncio
async def test_route_gate_verifies_retained_source_and_returns_actionable_error(
    monkeypatch,
):
    from app.routers import document_templates as router

    content = source("Fee: $400")
    read = AsyncMock(return_value=content)
    monkeypatch.setattr(router, "_verified_template_source", read)
    template = SimpleNamespace(format="docx")
    with pytest.raises(HTTPException) as error:
        await router._ensure_word_source_review(template, {"source_review_version": 1})
    assert error.value.status_code == 422
    read.assert_awaited_once_with(template)
    read.reset_mock()
    await router._ensure_word_source_review(template, {})
    await router._ensure_word_source_review(
        SimpleNamespace(format="pdf"), {"source_review_version": 1}
    )
    read.assert_not_awaited()


@pytest.mark.asyncio
async def test_publish_waits_for_source_decisions_even_after_a_successful_test(
    monkeypatch,
):
    from app.routers import document_templates as router
    from app.schemas.document_template import DocumentTemplatePublishRequest

    content = source("Retainer: $500")
    template = SimpleNamespace(
        id=uuid.uuid4(),
        format="docx",
        is_active=False,
        current_version_no=1,
        tested_version_no=1,
        published_version_no=None,
        variable_schema={"fields": [], "source_review_version": 1},
    )
    user = SimpleNamespace(id=uuid.uuid4(), tenant_id=uuid.uuid4())
    db = AsyncMock()
    db.scalar.return_value = template
    monkeypatch.setattr(router, "set_tenant_context", AsyncMock())
    monkeypatch.setattr(router, "record_version", AsyncMock())
    monkeypatch.setattr(
        router, "_verified_template_source", AsyncMock(return_value=content)
    )
    monkeypatch.setattr(
        router, "_template_response", lambda value: {"is_active": value.is_active}
    )
    with pytest.raises(HTTPException) as error:
        await router.publish_template(
            template.id, DocumentTemplatePublishRequest(), current_user=user, db=db
        )
    assert error.value.status_code == 422
    assert not template.is_active
    db.commit.assert_not_awaited()
    template.variable_schema["source_review"] = {
        item["id"]: "fixed" for item in docx_outline(content)["review_candidates"]
    }
    result = await router.publish_template(
        template.id, DocumentTemplatePublishRequest(), current_user=user, db=db
    )
    assert result == {"is_active": True}
    db.commit.assert_awaited_once()


def test_source_review_ignores_headings_and_invalid_calendar_dates():
    paragraph = {
        "ordinal": 0,
        "text": "Fees: section 14-05-09; due 12/25/26",
        "runs": [{"start": 0, "text": "Fees:", "bold": True}],
    }
    candidates, _ = source_review_candidates([paragraph])
    assert [item["source_text"] for item in candidates] == ["12/25/26"]


@pytest.mark.asyncio
async def test_outline_route_serializes_source_order_review_and_formatting(monkeypatch):
    from app.routers import document_templates as router

    def build(doc):
        doc.add_paragraph("Retainer: $500", style="List Number")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Signature: ___"

    content = source(build=build)
    template = SimpleNamespace(id=uuid.uuid4(), format="docx")
    user = SimpleNamespace(tenant_id=uuid.uuid4())
    db = AsyncMock()
    db.scalar.return_value = template
    monkeypatch.setattr(router, "set_tenant_context", AsyncMock())
    monkeypatch.setattr(
        router, "_verified_template_source", AsyncMock(return_value=content)
    )
    response = await router.get_template_outline(template.id, current_user=user, db=db)
    payload = response.model_dump(mode="json")
    assert [block["kind"] for block in payload["blocks"]] == ["paragraph", "table"]
    assert payload["paragraphs"][0]["numbering"] == "1."
    assert payload["paragraphs"][0]["alignment"] == "left"
    assert payload["paragraphs"][0]["dynamic_field"] is False
    assert [item["source_text"] for item in payload["review_candidates"]] == [
        "$500",
        "___",
    ]
    assert payload["review_truncated"] is False


def test_intake_review_metadata_and_choice_authority_cannot_be_removed():
    from app.routers.document_templates import _reviewed_variable_schema

    content = source("Answer: ___ yes ___ no")
    analysis = analyze_template_upload(
        filename="synthetic.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_bytes=content,
    )
    schema = analysis.variable_schema
    assert schema["source_review_version"] == 1
    submitted = {
        **schema,
        "source_review_version": 0,
        "source_review": {"fake": "fixed"},
        "fields": [
            {**field, "docx_choice": {}, "field_type": "text"}
            for field in schema["fields"]
        ],
    }
    reviewed = _reviewed_variable_schema(json.dumps(submitted), schema)
    assert reviewed["source_review_version"] == 1
    assert "source_review" not in reviewed
    assert all(
        field["docx_choice"]["option"] in {"Yes", "No"}
        and field["field_type"] == "checkbox"
        for field in reviewed["fields"]
    )
