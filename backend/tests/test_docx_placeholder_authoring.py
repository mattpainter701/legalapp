"""Regression tests for the derived Word placeholder source."""

import hashlib
import io
import uuid
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from docx import Document

from app.services.docx_outline import docx_outline
from app.services.docx_placeholder_authoring import (
    cleanup_docx_source,
    derived_source_is_current,
    derive_reviewed_docx_source,
    resolve_source_mode,
    suggest_source_mode,
)
from app.services.docx_templates import TemplateDocxError, iter_docx_paragraphs


def _source(build=None):
    doc = Document()
    doc.add_paragraph("Car value: [AMOUNT]; loan: [AMOUNT]")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Client: "
    table.cell(0, 0).paragraphs[0].add_run("Taylor Example").bold = True
    doc.sections[0].footer.paragraphs[0].text = "Date: [DATE]"
    if build:
        build(doc)
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_derivation_rewrites_only_explicitly_mapped_spans_and_preserves_evidence():
    content = _source()
    outline = docx_outline(content)
    amount_candidates = [
        item
        for item in outline["review_candidates"]
        if item["source_text"] == "[AMOUNT]"
    ]
    client = next(
        item
        for item in outline["review_candidates"]
        if item["source_text"] == "Taylor Example"
    )
    # Map only the first amount and the explicitly reviewed identity.  The
    # second amount and footer date remain source evidence for later review.
    result = derive_reviewed_docx_source(
        content,
        fields=[
            {
                "name": "car_value",
                "source_text": "[AMOUNT]",
                "docx_anchor": amount_candidates[0]["docx_anchor"],
            },
            {
                "name": "client_name",
                "source_text": client["source_text"],
                "docx_anchor": client["docx_anchor"],
            },
        ],
        source_mode="prose",
    )
    texts = [p.text for p in iter_docx_paragraphs(Document(io.BytesIO(result.content)))]
    assert "Car value: {{car_value}}; loan: [AMOUNT]" in texts
    assert any("{{client_name}}" in text for text in texts)
    assert any("[DATE]" in text for text in texts)
    assert result.metadata["original_sha256"] == hashlib.sha256(content).hexdigest()
    assert (
        result.metadata["derived_sha256"] == hashlib.sha256(result.content).hexdigest()
    )
    active_field = next(
        item for item in result.variable_schema["fields"] if item["name"] == "car_value"
    )
    assert active_field["source_text"] == "{{car_value}}"
    assert "docx_anchor" not in active_field
    assert result.variable_schema["source_review"] == {}
    assert derived_source_is_current(content, result.content, result.metadata)
    assert not derived_source_is_current(
        content + b"x", result.content, result.metadata
    )


def test_repeated_generic_spans_are_independent_and_formatting_split_runs_is_safe():
    content = _source()
    outline = docx_outline(content)
    amounts = [
        item
        for item in outline["review_candidates"]
        if item["source_text"] == "[AMOUNT]"
    ]
    result = derive_reviewed_docx_source(
        content,
        fields=[
            {
                "name": "car_value",
                "source_text": "[AMOUNT]",
                "docx_anchor": amounts[0]["docx_anchor"],
            },
            {
                "name": "loan_value",
                "source_text": "[AMOUNT]",
                "docx_anchor": amounts[1]["docx_anchor"],
            },
        ],
    )
    text = "\n".join(
        p.text for p in iter_docx_paragraphs(Document(io.BytesIO(result.content)))
    )
    assert "Car value: {{car_value}}; loan: {{loan_value}}" in text


def test_unanchored_inferred_identity_is_not_silently_rewritten():
    content = _source()
    result = derive_reviewed_docx_source(
        content,
        fields=[{"name": "client_name", "source_text": "Taylor Example"}],
    )
    assert "Taylor Example" in "\n".join(
        p.text for p in iter_docx_paragraphs(Document(io.BytesIO(result.content)))
    )


def test_explicit_arbitrary_prose_selection_is_allowed_when_anchor_matches():
    content = _source()
    source = "Car value"
    result = derive_reviewed_docx_source(
        content,
        fields=[
            {
                "name": "caption",
                "source_text": source,
                "docx_anchor": {"paragraph_ordinal": 0, "start": 0, "end": len(source)},
            }
        ],
    )
    assert "{{caption}}: [AMOUNT]" in "\n".join(
        p.text for p in iter_docx_paragraphs(Document(io.BytesIO(result.content)))
    )


def test_mode_is_a_suggestion_and_invalid_override_is_rejected():
    content = _source()
    suggestion = suggest_source_mode(content)
    assert suggestion.suggested_mode == "prose"
    assert resolve_source_mode(suggestion) == "prose"
    assert resolve_source_mode(suggestion, "form") == "form"
    with pytest.raises(TemplateDocxError, match="prose or form"):
        resolve_source_mode(suggestion, "unknown")


def test_conflicting_disposition_and_overlapping_mapping_fail_closed():
    content = _source()
    candidate = next(
        item
        for item in docx_outline(content)["review_candidates"]
        if item["source_text"] == "[AMOUNT]"
    )
    field = {
        "name": "amount",
        "source_text": "[AMOUNT]",
        "docx_anchor": candidate["docx_anchor"],
    }
    with pytest.raises(TemplateDocxError, match="conflicts"):
        derive_reviewed_docx_source(
            content, fields=[field], decisions={candidate["id"]: "fixed"}
        )
    with pytest.raises(TemplateDocxError, match="same source span"):
        derive_reviewed_docx_source(content, fields=[field, {**field, "name": "other"}])


def test_signature_disposition_requires_signature_compatible_field():
    content = _source()
    candidate = next(
        item
        for item in docx_outline(content)["review_candidates"]
        if item["source_text"] == "[AMOUNT]"
    )
    field = {
        "name": "signer",
        "source_text": "[AMOUNT]",
        "docx_anchor": candidate["docx_anchor"],
    }
    with pytest.raises(TemplateDocxError, match="conflicts"):
        derive_reviewed_docx_source(
            content, fields=[field], decisions={candidate["id"]: "signature"}
        )
    result = derive_reviewed_docx_source(
        content,
        fields=[{**field, "field_type": "signature"}],
        decisions={candidate["id"]: "signature"},
    )
    assert "{{signer}}" in "\n".join(
        p.text for p in iter_docx_paragraphs(Document(io.BytesIO(result.content)))
    )


def test_cleanup_preserves_tokens_and_rejects_token_edits():
    doc = Document()
    doc.add_paragraph("Clause {{client_name}} stray")
    stream = io.BytesIO()
    doc.save(stream)
    content = stream.getvalue()
    cleaned = cleanup_docx_source(
        content,
        paragraph_ordinal=0,
        start=23,
        end=28,
        original_text="stray",
        replacement_text="text",
    )
    assert "Clause {{client_name}} text" in "\n".join(
        p.text for p in iter_docx_paragraphs(Document(io.BytesIO(cleaned)))
    )
    with pytest.raises(TemplateDocxError, match="placeholder tokens"):
        cleanup_docx_source(
            content,
            paragraph_ordinal=0,
            start=7,
            end=22,
            original_text="{{client_name}}",
            replacement_text="removed",
        )


@pytest.mark.asyncio
async def test_derive_word_draft_creates_new_inactive_source_owned_by_server(
    monkeypatch,
):
    from app.routers import document_templates as router
    from app.schemas.document_template import DocumentTemplateWordDeriveRequest

    content = _source()
    outline = docx_outline(content)
    candidate = next(
        item
        for item in outline["review_candidates"]
        if item["source_text"] == "[AMOUNT]"
    )
    original_id = uuid.uuid4()
    original = SimpleNamespace(
        id=original_id,
        tenant_id=uuid.UUID("11111111-1111-1111-1111-111111111111"),
        format="docx",
        source_storage_path="/tmp/original.docx",
        source_filename="master.docx",
        source_content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        source_sha256=hashlib.sha256(content).hexdigest(),
        title="Master",
        body="",
        category="other",
        description=None,
        visibility="tenant",
        layer=None,
        module=None,
        stage=None,
        jurisdiction=None,
        kind=None,
        variable_schema={"source_review": {}},
        signer_roles=None,
        branding_profile=None,
    )
    db = AsyncMock()
    db.add = Mock()
    db.scalar.return_value = original
    monkeypatch.setattr(router, "set_tenant_context", AsyncMock())
    monkeypatch.setattr(
        router, "_verified_template_source", AsyncMock(return_value=content)
    )
    persisted = []

    async def persist(**kwargs):
        persisted.append(kwargs)
        return "F:/derived.docx"

    monkeypatch.setattr(router, "_persist_template_source", persist)
    monkeypatch.setattr(router, "_template_response", lambda value: value)
    user = SimpleNamespace(tenant_id=original.tenant_id)
    response = await router.derive_word_draft(
        original_id,
        DocumentTemplateWordDeriveRequest(
            fields=[
                {
                    "name": "amount",
                    "source_text": "[AMOUNT]",
                    "docx_anchor": candidate["docx_anchor"],
                }
            ]
        ),
        current_user=user,
        db=db,
    )
    assert response.is_active is False
    assert response.published_version_no is None
    assert response.variable_schema["source"] == "docx_derived_placeholder"
    assert persisted[0]["content"] != content
    db.commit.assert_awaited_once()


def test_failed_draft_save_removes_only_request_owned_files(tmp_path):
    from app.routers.document_templates import _remove_created_template_files

    owned = tmp_path / "derived.docx"
    retained = tmp_path / "original.docx"
    unrelated = tmp_path / "customer-upload.docx"
    owned.write_bytes(b"derived")
    retained.write_bytes(b"retained")
    unrelated.write_bytes(b"keep")

    _remove_created_template_files([str(owned), str(retained)])

    assert not owned.exists()
    assert not retained.exists()
    assert unrelated.read_bytes() == b"keep"


def _route_template(content, *, evidence_path=None):
    tenant = uuid.UUID("11111111-1111-1111-1111-111111111111")
    return SimpleNamespace(
        id=uuid.uuid4(),
        tenant_id=tenant,
        format="docx",
        source_storage_path="master.docx",
        source_filename="master.docx",
        source_content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        source_sha256=hashlib.sha256(content).hexdigest(),
        title="Master",
        body="",
        category="other",
        description=None,
        visibility="tenant",
        layer=None,
        module=None,
        stage=None,
        jurisdiction=None,
        kind=None,
        variable_schema={"source_review": {}},
        signer_roles=None,
        branding_profile=None,
        source_evidence_storage_path=str(evidence_path) if evidence_path else None,
        source_evidence_filename="original.docx" if evidence_path else None,
        source_evidence_sha256=hashlib.sha256(evidence_path.read_bytes()).hexdigest()
        if evidence_path
        else None,
    )


@pytest.mark.asyncio
async def test_derive_second_storage_failure_removes_first_file_and_keeps_original(
    tmp_path, monkeypatch
):
    from app.routers import document_templates as router
    from app.schemas.document_template import DocumentTemplateWordDeriveRequest

    content = _source()
    template = _route_template(content)
    candidate = next(
        item
        for item in docx_outline(content)["review_candidates"]
        if item["source_text"] == "[AMOUNT]"
    )
    first = tmp_path / "derived.docx"
    original = tmp_path / "master.docx"
    original.write_bytes(content)
    calls = 0

    async def persist(**kwargs):
        nonlocal calls
        calls += 1
        first.write_bytes(kwargs["content"])
        if calls == 2:
            raise OSError("evidence write failed")
        return str(first)

    db = AsyncMock()
    db.add = Mock()
    db.scalar.return_value = template
    monkeypatch.setattr(router, "set_tenant_context", AsyncMock())
    monkeypatch.setattr(
        router, "_verified_template_source", AsyncMock(return_value=content)
    )
    monkeypatch.setattr(router, "_persist_template_source", persist)
    with pytest.raises(Exception, match="source could not be saved"):
        await router.derive_word_draft(
            template.id,
            DocumentTemplateWordDeriveRequest(
                fields=[
                    {
                        "name": "amount",
                        "source_text": "[AMOUNT]",
                        "docx_anchor": candidate["docx_anchor"],
                    }
                ]
            ),
            current_user=SimpleNamespace(tenant_id=template.tenant_id),
            db=db,
        )
    assert not first.exists()
    assert original.read_bytes() == content


@pytest.mark.asyncio
@pytest.mark.parametrize("rollback_fails", [False, True])
async def test_cleanup_commit_failure_removes_new_files_and_keeps_original(
    tmp_path, monkeypatch, rollback_fails
):
    from app.routers import document_templates as router
    from app.schemas.document_template import DocumentTemplateWordCleanupRequest

    content = _source()
    evidence = tmp_path / "original.docx"
    evidence.write_bytes(content)
    template = _route_template(content, evidence_path=evidence)
    monkeypatch.setattr(router, "_template_source_dir", lambda *_: str(tmp_path))
    paths = []

    async def persist(**kwargs):
        path = tmp_path / f"new-{len(paths)}.docx"
        path.write_bytes(kwargs["content"])
        paths.append(path)
        return str(path)

    db = AsyncMock()
    db.add = Mock()
    db.scalar.return_value = template
    db.commit.side_effect = RuntimeError("db failed")
    if rollback_fails:
        db.rollback.side_effect = RuntimeError("connection lost")
    monkeypatch.setattr(router, "set_tenant_context", AsyncMock())
    monkeypatch.setattr(
        router, "_verified_template_source", AsyncMock(return_value=content)
    )
    monkeypatch.setattr(router, "_persist_template_source", persist)
    payload = DocumentTemplateWordCleanupRequest(
        paragraph_ordinal=0,
        start=0,
        end=9,
        original_text="Car value",
        replacement_text="Car value",
    )
    with pytest.raises(Exception, match="cleaned Word draft could not be saved"):
        await router.cleanup_word_draft(
            template.id,
            payload,
            current_user=SimpleNamespace(tenant_id=template.tenant_id),
            db=db,
        )
    assert all(not path.exists() for path in paths)
    assert evidence.read_bytes() == content
