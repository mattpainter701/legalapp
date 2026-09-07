"""Regression coverage for tenant-scoped original Word evidence."""

import hashlib
import io
import uuid
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from fastapi import HTTPException
from docx import Document

from app.services.docx_outline import docx_outline


TENANT = uuid.UUID("11111111-1111-1111-1111-111111111111")
OTHER_TENANT = uuid.UUID("22222222-2222-2222-2222-222222222222")


def _docx(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _template(*, tenant_id=TENANT, template_id=None, source=b"master", **overrides):
    template_id = template_id or uuid.uuid4()
    values = {
        "id": template_id,
        "tenant_id": tenant_id,
        "format": "docx",
        "source_storage_path": "master.docx",
        "source_filename": "master.docx",
        "source_content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "source_sha256": hashlib.sha256(source).hexdigest(),
        "source_evidence_storage_path": None,
        "source_evidence_filename": None,
        "source_evidence_sha256": None,
        "source_provenance": None,
    }
    values.update(overrides)
    return SimpleNamespace(**values)


@pytest.mark.asyncio
async def test_verified_word_original_falls_back_to_verified_master(monkeypatch):
    from app.routers import document_templates as router

    content = _docx("Original master")
    template = _template(source=content)
    verified = AsyncMock(return_value=content)
    monkeypatch.setattr(router, "_verified_template_source", verified)

    original, filename = await router._verified_word_original(template)

    assert original == content
    assert filename == "master.docx"
    verified.assert_awaited_once_with(template)


@pytest.mark.asyncio
async def test_verified_word_original_reads_oldest_evidence_with_digest_and_tenant_scope(
    tmp_path, monkeypatch
):
    from app.routers import document_templates as router

    content = _docx("Original upload")
    template_id = uuid.uuid4()
    evidence_root = tmp_path / str(TENANT) / str(template_id)
    evidence_root.mkdir(parents=True)
    evidence = evidence_root / "original-master.docx"
    evidence.write_bytes(content)
    template = _template(
        template_id=template_id,
        source=b"derived source",
        source_evidence_storage_path=str(evidence),
        source_evidence_filename="original-master.docx",
        source_evidence_sha256=hashlib.sha256(content).hexdigest(),
        source_provenance={
            "kind": "derived",
            "original_template_id": str(uuid.uuid4()),
            "original_sha256": hashlib.sha256(content).hexdigest(),
        },
    )
    monkeypatch.setattr(
        router,
        "_template_source_dir",
        lambda tenant_id, _id: str(tmp_path / tenant_id / str(_id)),
    )

    original, filename = await router._verified_word_original(template)

    assert original == content
    assert filename == "original-master.docx"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "changes",
    [
        {"source_evidence_storage_path": None},
        {"source_evidence_sha256": None},
        {"source_evidence_sha256": "0" * 64},
        {
            "source_evidence_storage_path": None,
            "source_evidence_sha256": None,
            "source_provenance": {"kind": "derived"},
        },
        {"source_evidence_storage_path": "__foreign_tenant__"},
    ],
)
async def test_verified_word_original_rejects_incomplete_or_tampered_derived_evidence(
    tmp_path, monkeypatch, changes
):
    from app.routers import document_templates as router

    template_id = uuid.uuid4()
    root = tmp_path / str(TENANT) / str(template_id)
    root.mkdir(parents=True)
    evidence = root / "original.docx"
    evidence_content = _docx("retained original")
    evidence.write_bytes(evidence_content)
    values = {
        "template_id": template_id,
        "source": b"derived",
        "source_evidence_storage_path": str(evidence),
        "source_evidence_filename": "original.docx",
        "source_evidence_sha256": hashlib.sha256(evidence_content).hexdigest(),
        "source_provenance": {
            "kind": "derived",
            "original_template_id": str(uuid.uuid4()),
        },
    }
    values.update(changes)
    if values["source_evidence_storage_path"] == "__foreign_tenant__":
        foreign_root = tmp_path / str(OTHER_TENANT) / str(template_id)
        foreign_root.mkdir(parents=True)
        foreign_evidence = foreign_root / "original.docx"
        foreign_evidence.write_bytes(evidence_content)
        values["source_evidence_storage_path"] = str(foreign_evidence)
    template = _template(**values)
    monkeypatch.setattr(
        router,
        "_template_source_dir",
        lambda tenant_id, _id: str(tmp_path / tenant_id / str(_id)),
    )
    verified = AsyncMock()
    monkeypatch.setattr(router, "_verified_template_source", verified)

    with pytest.raises(HTTPException) as error:
        await router._verified_word_original(template)

    assert error.value.status_code == 409
    verified.assert_not_awaited()


@pytest.mark.asyncio
async def test_second_derive_retains_root_evidence_and_uses_real_tenant_storage(
    tmp_path, monkeypatch
):
    from app.routers import document_templates as router
    from app.schemas.document_template import DocumentTemplateWordDeriveRequest

    parent_content = _docx("Amount: [AMOUNT]")
    original = _docx("Root original amount")
    root_id = uuid.uuid4()
    parent_id = uuid.uuid4()
    upload_root = tmp_path / "uploads"
    monkeypatch.setattr(router.settings, "UPLOAD_DIR", str(upload_root))
    parent_dir = upload_root / str(TENANT) / "templates" / str(parent_id)
    parent_dir.mkdir(parents=True)
    parent_source = parent_dir / "parent.docx"
    parent_source.write_bytes(parent_content)
    retained = parent_dir / "original-parent.docx"
    retained.write_bytes(original)
    parent = _template(
        template_id=parent_id,
        source=parent_content,
        source_storage_path=str(parent_source),
        source_filename="parent.docx",
        source_evidence_storage_path=str(retained),
        source_evidence_filename="original-parent.docx",
        source_evidence_sha256=hashlib.sha256(original).hexdigest(),
        source_provenance={
            "kind": "derived",
            "original_template_id": str(root_id),
            "original_sha256": hashlib.sha256(original).hexdigest(),
        },
        title="Parent",
        body="",
        category="other",
        description=None,
        visibility="tenant",
        layer=None,
        module=None,
        stage=None,
        jurisdiction=None,
        kind=None,
        variable_schema={},
        signer_roles=None,
        branding_profile=None,
    )
    candidate = next(
        item
        for item in docx_outline(parent_content)["review_candidates"]
        if item["source_text"] == "[AMOUNT]"
    )
    db = AsyncMock()
    db.scalar.return_value = parent
    captured = []
    db.add = Mock(side_effect=captured.append)
    monkeypatch.setattr(router, "set_tenant_context", AsyncMock())
    monkeypatch.setattr(router, "_template_response", lambda value: value)

    response = await router.derive_word_draft(
        parent_id,
        DocumentTemplateWordDeriveRequest(
            fields=[
                {
                    "name": "amount",
                    "source_text": "[AMOUNT]",
                    "docx_anchor": candidate["docx_anchor"],
                }
            ]
        ),
        current_user=SimpleNamespace(tenant_id=TENANT),
        db=db,
    )

    child = captured[0]
    assert response is child
    active_content = Path(child.source_storage_path).read_bytes()
    assert active_content != parent_content
    assert any(
        "{{amount}}" in paragraph.text
        for paragraph in Document(io.BytesIO(active_content)).paragraphs
    )
    assert any(
        field.get("name") == "amount" for field in child.variable_schema["fields"]
    )
    assert Path(child.source_evidence_storage_path).read_bytes() == original
    assert child.source_evidence_sha256 == hashlib.sha256(original).hexdigest()
    assert child.source_evidence_filename == "original-parent.docx"
    assert child.source_provenance["original_template_id"] == str(root_id)
    assert child.source_provenance["parent_template_id"] == str(parent_id)
    assert parent_source.read_bytes() == parent_content
    assert retained.read_bytes() == original


@pytest.mark.asyncio
async def test_cleanup_retains_root_evidence_and_clears_source_review(
    tmp_path, monkeypatch
):
    from app.routers import document_templates as router
    from app.schemas.document_template import DocumentTemplateWordCleanupRequest

    active_content = _docx("Fee: $400")
    original = _docx("Root original fee")
    root_id = uuid.uuid4()
    parent_id = uuid.uuid4()
    upload_root = tmp_path / "uploads"
    monkeypatch.setattr(router.settings, "UPLOAD_DIR", str(upload_root))
    parent_dir = upload_root / str(TENANT) / "templates" / str(parent_id)
    parent_dir.mkdir(parents=True)
    parent_source = parent_dir / "parent.docx"
    parent_source.write_bytes(active_content)
    retained = parent_dir / "original-parent.docx"
    retained.write_bytes(original)
    parent = _template(
        template_id=parent_id,
        source=active_content,
        source_storage_path=str(parent_source),
        source_filename="parent.docx",
        source_evidence_storage_path=str(retained),
        source_evidence_filename="original-parent.docx",
        source_evidence_sha256=hashlib.sha256(original).hexdigest(),
        source_provenance={"original_template_id": str(root_id)},
        variable_schema={"source_review": {"candidate": "fixed"}},
        title="Parent",
        body="",
        category="other",
        description=None,
        visibility="tenant",
        layer=None,
        module=None,
        stage=None,
        jurisdiction=None,
        kind=None,
        signer_roles=None,
        branding_profile=None,
    )
    db = AsyncMock()
    db.scalar.return_value = parent
    captured = []
    db.add = Mock(side_effect=captured.append)
    monkeypatch.setattr(router, "set_tenant_context", AsyncMock())
    monkeypatch.setattr(router, "_template_response", lambda value: value)

    response = await router.cleanup_word_draft(
        parent_id,
        DocumentTemplateWordCleanupRequest(
            paragraph_ordinal=0,
            start=5,
            end=9,
            original_text="$400",
            replacement_text="$450",
        ),
        current_user=SimpleNamespace(tenant_id=TENANT),
        db=db,
    )

    child = captured[0]
    assert response is child
    assert child.variable_schema["source_review"] == {}
    child_content = Path(child.source_storage_path).read_bytes()
    assert any(
        "Fee: $450" in paragraph.text
        for paragraph in Document(io.BytesIO(child_content)).paragraphs
    )
    assert Path(child.source_evidence_storage_path).read_bytes() == original
    assert child.source_evidence_filename == "original-parent.docx"
    assert child.source_provenance["original_template_id"] == str(root_id)
    assert child.source_provenance["parent_template_id"] == str(parent_id)
    assert parent_source.read_bytes() == active_content


@pytest.mark.asyncio
async def test_original_source_route_returns_master_for_base_template(monkeypatch):
    from app.routers import document_templates as router

    content = _docx("Base source")
    template = _template(source=content)
    db = AsyncMock()
    db.scalar.return_value = template
    monkeypatch.setattr(router, "set_tenant_context", AsyncMock())
    monkeypatch.setattr(
        router,
        "_verified_word_original",
        AsyncMock(return_value=(content, "master.docx")),
    )

    response = await router.download_original_template_source(
        template.id,
        current_user=SimpleNamespace(tenant_id=TENANT),
        db=db,
    )

    assert response.body == content
    assert "master.docx" in response.headers["content-disposition"]


@pytest.mark.asyncio
async def test_original_source_route_returns_retained_evidence_for_derived_template(
    monkeypatch,
):
    from app.routers import document_templates as router

    content = _docx("Oldest source")
    template = _template(
        source=b"derived source",
        source_provenance={
            "kind": "derived",
            "original_template_id": str(uuid.uuid4()),
        },
    )
    db = AsyncMock()
    db.scalar.return_value = template
    monkeypatch.setattr(router, "set_tenant_context", AsyncMock())
    monkeypatch.setattr(
        router,
        "_verified_word_original",
        AsyncMock(return_value=(content, "original-master.docx")),
    )

    response = await router.download_original_template_source(
        template.id,
        current_user=SimpleNamespace(tenant_id=TENANT),
        db=db,
    )

    assert response.body == content
    assert "original-master.docx" in response.headers["content-disposition"]


@pytest.mark.asyncio
async def test_original_source_route_is_tenant_scoped_and_hides_missing_templates(
    monkeypatch,
):
    from app.routers import document_templates as router

    template_id = uuid.uuid4()
    db = AsyncMock()
    db.scalar.return_value = None
    monkeypatch.setattr(router, "set_tenant_context", AsyncMock())

    with pytest.raises(HTTPException) as error:
        await router.download_original_template_source(
            template_id,
            current_user=SimpleNamespace(tenant_id=OTHER_TENANT),
            db=db,
        )

    assert error.value.status_code == 404
    query = db.scalar.await_args.args[0]
    assert any(value == OTHER_TENANT for value in query.compile().params.values())
