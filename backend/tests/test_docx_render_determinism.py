"""Preview and save refill Word templates at different wall-clock times."""

import io
import zipfile
import shutil

import pytest

from docx import Document

from app.services.docx_templates import fill_docx_template
from app.services.docx_to_pdf import docx_to_pdf_bytes


def test_repeated_fill_preserves_identical_bytes_across_archive_timestamps(monkeypatch):
    document = Document()
    document.add_paragraph("Matter: {{matter_name}}")
    document.sections[0].footer.paragraphs[0].text = "Internal QA"
    source = io.BytesIO()
    document.save(source)
    schema = {"fields": [{"name": "matter_name", "type": "text", "required": True}]}

    def render(hour, name="Synthetic QA Matter"):
        monkeypatch.setattr(
            zipfile.time,
            "localtime",
            lambda *args: (2026, 9, 8, hour, 0, 0, 1, 251, 0),
        )
        return fill_docx_template(
            source.getvalue(),
            variable_schema=schema,
            variables={"matter_name": name},
            enforce_required=True,
        )

    first = render(12)
    second = render(14)
    assert first == second
    assert render(15, "Different Matter") != first
    reopened = Document(io.BytesIO(second))
    assert reopened.paragraphs[0].text == "Matter: Synthetic QA Matter"
    assert reopened.sections[0].footer.paragraphs[0].text == "Internal QA"


@pytest.mark.asyncio
@pytest.mark.skipif(not shutil.which("libreoffice"), reason="LibreOffice not installed")
async def test_real_converter_preserves_reviewed_pdf_across_independent_fills(
    monkeypatch,
):
    document = Document()
    document.add_paragraph("Matter: {{matter_name}}")
    source = io.BytesIO()
    document.save(source)
    outputs = []
    for hour, name in ((12, "Cedar QA"), (14, "Cedar QA"), (15, "Apex QA")):
        monkeypatch.setattr(
            zipfile.time,
            "localtime",
            lambda *args, hour=hour: (2026, 9, 8, hour, 0, 0, 1, 251, 0),
        )
        filled = fill_docx_template(
            source.getvalue(),
            variable_schema={"fields": [{"name": "matter_name", "type": "text"}]},
            variables={"matter_name": name},
        )
        outputs.append(
            await docx_to_pdf_bytes(filled, executable=shutil.which("libreoffice"))
        )
    assert outputs[0] == outputs[1]
    assert outputs[0] != outputs[2]
